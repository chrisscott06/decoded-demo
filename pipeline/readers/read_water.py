"""Water reader.

Reads the Dashboard sheet from the Westbrook Water workbook into per-site water
consumption + status flags. Output is keyed by canonical site ID and contains
all 13 canonical sites — sites without water data get data_status="missing".
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from site_resolver import CANONICAL_SITES, is_known_unmapped, resolve

# Map source data_quality strings → DataStatusBadge state
_QUALITY_TO_STATUS = {
    "actual": "confirmed",
    "estimated": "partial",
    "": "missing",
    "no data": "missing",
    "none": "missing",
    "tbc": "missing",
    "unknown": "missing",
}


def _to_str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s == "—":  # em-dash placeholder
        return None
    return s


def _to_int(v: Any) -> int | None:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == "" or s == "—" or s.lower() in {"unknown", "tbc", "n/a", "none"}:
            return None
        try:
            return int(round(float(s)))
        except ValueError:
            return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def _quality_to_status(quality: str | None, log: list[str], site_id: str) -> str:
    if quality is None:
        return "missing"
    key = quality.strip().lower()
    if key in _QUALITY_TO_STATUS:
        return _QUALITY_TO_STATUS[key]
    log.append(f"  Water: unrecognised data_quality {quality!r} for {site_id} — defaulting to 'partial'")
    return "partial"


def find_workbook(source_data: Path) -> Path:
    candidates = list(source_data.glob("*CA-X-1002*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No Water workbook (CA-X-1002*) in {source_data}")
    return candidates[0]


def read_water(workbook_path: Path, log: list[str]) -> dict[str, dict]:
    log.append("")
    log.append(f"[Water] Reading: {workbook_path.name}")

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    if "Dashboard" not in wb.sheetnames:
        log.append(f"  ERROR: 'Dashboard' sheet missing. Available: {wb.sheetnames}")
        return {}

    ws = wb["Dashboard"]

    out: dict[str, dict] = {}
    rows_scanned = 0
    rows_resolved = 0
    unmapped: list[str] = []

    # Data rows 5-17 (13 sites), then row 18 PORTFOLIO TOTAL, then summary/notes (skip).
    # We stop scanning the moment we hit PORTFOLIO TOTAL — anything below is not site data.
    for r in range(5, ws.max_row + 1):
        site_name = _to_str(ws.cell(r, 1).value)
        if site_name is None:
            continue
        if is_known_unmapped(site_name):
            # PORTFOLIO TOTAL = end of site data. Don't scan further.
            break
        rows_scanned += 1
        site_id = resolve(site_name)
        if site_id is None:
            unmapped.append(site_name)
            continue

        quality = _to_str(ws.cell(r, 8).value)
        status = _quality_to_status(quality, log, site_id)

        out[site_id] = {
            "id": site_id,
            "water_company": _to_str(ws.cell(r, 2).value),
            "retailer": _to_str(ws.cell(r, 3).value),
            "meters_known": _to_int(ws.cell(r, 4).value),
            "meters_with_cy2025_data": _to_int(ws.cell(r, 5).value),
            "consumption_m3": _to_int(ws.cell(r, 6).value),
            "annual_estimate_m3": _to_int(ws.cell(r, 7).value),
            "data_quality": quality,
            "data_status": status,
            "cy2025_coverage": _to_str(ws.cell(r, 9).value),
            "completeness": _to_str(ws.cell(r, 10).value),
            "key_gaps": _to_str(ws.cell(r, 11).value),
            "arbnco_meters": _to_int(ws.cell(r, 12).value),
        }
        rows_resolved += 1

    # Ensure every canonical site has an entry, even if the workbook didn't list them
    missing_canonical = []
    for sid, info in CANONICAL_SITES.items():
        if sid not in out:
            missing_canonical.append(sid)
            out[sid] = {
                "id": sid,
                "water_company": None,
                "retailer": None,
                "meters_known": None,
                "meters_with_cy2025_data": None,
                "consumption_m3": None,
                "annual_estimate_m3": None,
                "data_quality": None,
                "data_status": "missing",
                "cy2025_coverage": None,
                "completeness": None,
                "key_gaps": None,
                "arbnco_meters": None,
            }

    log.append(f"  Dashboard: scanned={rows_scanned}, resolved={rows_resolved}, total in output={len(out)}")
    if unmapped:
        log.append(f"  Water: UNMAPPED names: {unmapped}")
    if missing_canonical:
        log.append(f"  Water: canonical sites not in workbook (filled as 'missing'): {missing_canonical}")

    return out
