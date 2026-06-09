"""Site Overview reader.

Reads four sheets from the Site Overview workbook and produces a per-site
record dict keyed by canonical site ID. See docs/briefs/ivg-data-source-inventory.md
for the canonical structure of each sheet.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from site_resolver import (
    CANONICAL_SITES,
    is_known_unmapped,
    resolve,
)


# ----- coercion helpers ---------------------------------------------------

def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def to_str(v: Any) -> str | None:
    if _is_blank(v):
        return None
    return str(v).strip()


def to_int(v: Any) -> int | None:
    if _is_blank(v):
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def to_float(v: Any) -> float | None:
    if _is_blank(v):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def to_pct(v: Any) -> int | None:
    """Coerce a percent value. Excel stores "73%" either as the string "73%" or
    as the float 0.73. Returns int 0-100 or None.
    """
    if _is_blank(v):
        return None
    if isinstance(v, str):
        s = v.strip().rstrip("%").strip()
        try:
            n = float(s)
        except ValueError:
            return None
        return int(round(n))
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    # 0-1 → assume fraction; >1 → assume already a percent
    if 0 <= n <= 1:
        return int(round(n * 100))
    return int(round(n))


def to_bool(v: Any) -> bool | None:
    if _is_blank(v):
        return None
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in {"yes", "y", "true", "1", "✓", "tick"}:
        return True
    if s in {"no", "n", "false", "0", "—", "-"}:
        return False
    return None


# ----- sheet readers ------------------------------------------------------

def _read_site_overview_sheet(ws, log: list[str]) -> dict[str, dict]:
    """Read Site Overview (sheet 1). Header row 4, data rows 5-17, row 18 = portfolio total (skip)."""
    out: dict[str, dict] = {}
    rows_scanned = 0
    rows_resolved = 0
    rows_skipped = 0
    unmapped: list[str] = []

    for r in range(5, ws.max_row + 1):
        site_name = to_str(ws.cell(r, 1).value)
        if site_name is None:
            rows_skipped += 1
            continue
        rows_scanned += 1
        if is_known_unmapped(site_name):
            rows_skipped += 1
            continue
        site_id = resolve(site_name)
        if site_id is None:
            unmapped.append(site_name)
            continue

        out[site_id] = {
            "id": site_id,
            "display_name": CANONICAL_SITES[site_id]["display"],
            "ref": to_str(ws.cell(r, 2).value),
            "identity": {
                "total_units": to_int(ws.cell(r, 3).value),
                "completed": to_int(ws.cell(r, 4).value),
                "void": to_int(ws.cell(r, 5).value),
                "occupancy_pct": to_pct(ws.cell(r, 6).value),
                "resi_gia_m2": to_float(ws.cell(r, 7).value),
                "landlord_gia_m2": to_float(ws.cell(r, 8).value),
                "total_gia_m2": to_float(ws.cell(r, 9).value),
                "years_built": to_str(ws.cell(r, 10).value),
                "total_phases": to_int(ws.cell(r, 11).value),
                "operational_phases": to_int(ws.cell(r, 12).value),
                "management_entity": to_str(ws.cell(r, 30).value),
            },
            "archetype": {
                "primary_heating": to_str(ws.cell(r, 13).value),
                "hot_water": to_str(ws.cell(r, 14).value),
                "chp": to_str(ws.cell(r, 15).value),
                "heat_network": to_str(ws.cell(r, 16).value),
                "heat_billing": to_str(ws.cell(r, 17).value),
                "grid_type": to_str(ws.cell(r, 18).value),
                "capacity_mva": to_float(ws.cell(r, 19).value),
                "asc_kva": to_float(ws.cell(r, 20).value),
                "solar_pv_kwp": to_float(ws.cell(r, 21).value),
                "battery": to_str(ws.cell(r, 22).value),
                "ev_chargers": to_str(ws.cell(r, 23).value),
                "pool_spa": to_str(ws.cell(r, 24).value),
                "restaurant": to_str(ws.cell(r, 25).value),
                "village_centre": to_str(ws.cell(r, 26).value),
                "sycous": to_str(ws.cell(r, 27).value),
                "ll_deduction": to_str(ws.cell(r, 28).value),
                "metering_arrangement": to_str(ws.cell(r, 29).value),
            },
            # Filled in by other sheet readers below (or stays null)
            "scope_allocation": None,
            "phasing": None,
            "energy_benchmarks": None,
            "notes": to_str(ws.cell(r, 31).value),
        }
        rows_resolved += 1

    log.append(f"  Site Overview: scanned={rows_scanned}, resolved={rows_resolved}, skipped={rows_skipped}")
    if unmapped:
        log.append(f"  Site Overview: UNMAPPED names: {unmapped}")
    return out


def _read_scope_allocation(ws, sites: dict[str, dict], log: list[str]) -> None:
    rows_scanned = 0
    rows_filled = 0
    unmapped: list[str] = []

    for r in range(5, ws.max_row + 1):
        site_name = to_str(ws.cell(r, 1).value)
        if site_name is None or is_known_unmapped(site_name):
            continue
        rows_scanned += 1
        site_id = resolve(site_name)
        if site_id is None:
            unmapped.append(site_name)
            continue
        if site_id not in sites:
            log.append(f"  Scope Allocation: site {site_id!r} not in Site Overview output — skipping")
            continue

        sites[site_id]["scope_allocation"] = {
            "gas_landlord": to_str(ws.cell(r, 3).value),
            "gas_occupied_resi": to_str(ws.cell(r, 4).value),
            "gas_void": to_str(ws.cell(r, 5).value),
            "elec_communal_ll": to_str(ws.cell(r, 6).value),
            "elec_occupied_resi": to_str(ws.cell(r, 7).value),
            "elec_void": to_str(ws.cell(r, 8).value),
            "has_gas": to_bool(ws.cell(r, 10).value),
            "sycous_deduction": to_bool(ws.cell(r, 11).value),
            "notes": to_str(ws.cell(r, 12).value),
        }
        rows_filled += 1

    log.append(f"  Scope Allocation: scanned={rows_scanned}, filled={rows_filled}")
    if unmapped:
        log.append(f"  Scope Allocation: UNMAPPED names: {unmapped}")


def _read_phasing(ws, sites: dict[str, dict], log: list[str]) -> None:
    rows_scanned = 0
    rows_filled = 0
    unmapped: list[str] = []

    for r in range(5, ws.max_row + 1):
        site_name = to_str(ws.cell(r, 1).value)
        if site_name is None or is_known_unmapped(site_name):
            continue
        rows_scanned += 1
        site_id = resolve(site_name)
        if site_id is None:
            unmapped.append(site_name)
            continue
        if site_id not in sites:
            continue

        phases = []
        for i in range(4):
            phases.append({
                "phase": i + 1,
                "units": to_int(ws.cell(r, 4 + i).value),
                "complete_year": to_int(ws.cell(r, 10 + i).value),
            })

        sites[site_id]["phasing"] = {
            "open_year": to_int(ws.cell(r, 3).value),
            "phases": phases,
            "total_units_planned": to_int(ws.cell(r, 8).value),
            "current_units": to_int(ws.cell(r, 9).value),
            "status": to_str(ws.cell(r, 15).value),
            "notes": to_str(ws.cell(r, 16).value),
        }
        rows_filled += 1

    log.append(f"  Phasing & Units: scanned={rows_scanned}, filled={rows_filled}")
    if unmapped:
        log.append(f"  Phasing & Units: UNMAPPED names: {unmapped}")


def _read_energy_benchmarks(ws, sites: dict[str, dict], log: list[str]) -> None:
    rows_scanned = 0
    rows_filled = 0
    unmapped: list[str] = []

    for r in range(5, ws.max_row + 1):
        site_name = to_str(ws.cell(r, 1).value)
        if site_name is None or is_known_unmapped(site_name):
            continue
        rows_scanned += 1
        site_id = resolve(site_name)
        if site_id is None:
            unmapped.append(site_name)
            continue
        if site_id not in sites:
            continue

        sites[site_id]["energy_benchmarks"] = {
            "actual_elec_kwh": to_float(ws.cell(r, 5).value),
            "actual_gas_kwh": to_float(ws.cell(r, 6).value),
            "actual_total_kwh": to_float(ws.cell(r, 7).value),
            "eui_total_kwh_m2": to_float(ws.cell(r, 8).value),
            "eui_elec_kwh_m2": to_float(ws.cell(r, 9).value),
            "eui_gas_kwh_m2": to_float(ws.cell(r, 10).value),
            "modelled_gas_eui_kwh_m2": to_float(ws.cell(r, 11).value),
            "gas_vs_model": to_float(ws.cell(r, 12).value),
            "tm54_benchmark": to_str(ws.cell(r, 13).value),
            "notes": to_str(ws.cell(r, 14).value),
        }
        rows_filled += 1

    log.append(f"  Energy Benchmarks: scanned={rows_scanned}, filled={rows_filled}")
    if unmapped:
        log.append(f"  Energy Benchmarks: UNMAPPED names: {unmapped}")


# ----- public entry point -------------------------------------------------

def read_site_overview(workbook_path: Path, log: list[str]) -> dict[str, dict]:
    log.append("")
    log.append(f"[Site Overview] Reading: {workbook_path.name}")
    mtime = datetime.fromtimestamp(workbook_path.stat().st_mtime, tz=timezone.utc)
    log.append(f"  Last modified: {mtime.isoformat()}")

    wb = load_workbook(workbook_path, data_only=True, read_only=True)

    expected_sheets = ["Site Overview", "Scope Allocation", "Phasing & Units", "Energy Benchmarks"]
    missing = [s for s in expected_sheets if s not in wb.sheetnames]
    if missing:
        log.append(f"  ERROR: missing expected sheets: {missing}")
        return {}

    sites = _read_site_overview_sheet(wb["Site Overview"], log)
    _read_scope_allocation(wb["Scope Allocation"], sites, log)
    _read_phasing(wb["Phasing & Units"], sites, log)
    _read_energy_benchmarks(wb["Energy Benchmarks"], sites, log)

    # Cross-validate Site Overview ↔ Scope Allocation count
    so_count = len(sites)
    sa_count = sum(1 for s in sites.values() if s["scope_allocation"] is not None)
    if so_count != sa_count:
        log.append(f"  WARNING: Site Overview ({so_count}) vs Scope Allocation ({sa_count}) row count mismatch")

    log.append(f"  Final: {len(sites)} sites in output")
    return sites


def find_workbook(source_data: Path) -> Path:
    """Locate Site Overview workbook by document reference prefix."""
    candidates = list(source_data.glob("*CA-X-2001*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No Site Overview workbook (CA-X-2001*) in {source_data}")
    return candidates[0]
