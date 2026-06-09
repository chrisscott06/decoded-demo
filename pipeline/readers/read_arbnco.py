"""Arbnco reader.

Reads the Raw Data sheet from the Arbnco master data workbook and produces
per-site consumption + emissions, plus a portfolio rollup.

Phase 0 note: the brief explicitly defers landlord/resident split until Sycous
data is available. We record meter-aggregate values and flag the pending split
in each site's `notes` field.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from site_resolver import is_known_unmapped, resolve


def _to_float(v: Any) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _to_int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def _to_str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def find_workbook(source_data: Path) -> Path:
    candidates = list(source_data.glob("*CA-X_1001*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No Arbnco workbook (CA-X_1001*) in {source_data}")
    return candidates[0]


def read_arbnco(workbook_path: Path, sites: dict[str, dict], log: list[str]) -> dict:
    """Read Arbnco Raw Data and return a dict keyed by site ID, plus a 'portfolio' key.

    `sites` is the output of read_site_overview — used to look up scope allocation
    for the per-site notes. If a site is missing its scope_allocation, we log it
    but still include the row.
    """
    log.append("")
    log.append(f"[Arbnco] Reading: {workbook_path.name}")

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    if "Raw Data" not in wb.sheetnames:
        log.append(f"  ERROR: 'Raw Data' sheet missing. Available: {wb.sheetnames}")
        return {}

    ws = wb["Raw Data"]

    # Real column layout (header row 1):
    #   col 1: blank/index, col 2: Asset (site name), col 3: Fund,
    #   col 4: # elec meters, col 5: # gas meters,
    #   col 6: total kWh, col 7: elec kWh, col 8: gas kWh,
    #   col 9: total actual tCO2e, col 10: elec actual, col 11: gas actual,
    #   col 12: total nat avg tCO2e, col 13: elec nat avg, col 14: gas nat avg,
    #   col 15: active for updates, col 16: notes
    out: dict[str, dict] = {}
    portfolio_row: dict | None = None
    rows_scanned = 0
    rows_resolved = 0
    unmapped_known: list[str] = []
    unmapped_unknown: list[str] = []

    for r in range(2, ws.max_row + 1):
        asset = _to_str(ws.cell(r, 2).value)
        fund = _to_str(ws.cell(r, 3).value)

        # Portfolio total: Asset blank, Fund "Inspired Villages"
        if asset is None and fund and fund.lower() == "inspired villages":
            portfolio_row = {
                "total_electricity_meters": _to_int(ws.cell(r, 4).value),
                "total_gas_meters": _to_int(ws.cell(r, 5).value),
                "total_consumption_kwh": _to_float(ws.cell(r, 6).value),
                "total_electricity_kwh": _to_float(ws.cell(r, 7).value),
                "total_gas_kwh": _to_float(ws.cell(r, 8).value),
                "total_emissions_actual_tco2e": _to_float(ws.cell(r, 9).value),
                "total_electricity_actual_tco2e": _to_float(ws.cell(r, 10).value),
                "total_gas_actual_tco2e": _to_float(ws.cell(r, 11).value),
                "total_emissions_national_avg_tco2e": _to_float(ws.cell(r, 12).value),
                "total_electricity_national_avg_tco2e": _to_float(ws.cell(r, 13).value),
                "total_gas_national_avg_tco2e": _to_float(ws.cell(r, 14).value),
            }
            continue

        if asset is None:
            continue

        rows_scanned += 1

        if is_known_unmapped(asset):
            unmapped_known.append(asset)
            continue

        site_id = resolve(asset)
        if site_id is None:
            unmapped_unknown.append(asset)
            continue

        # Scope allocation lookup (informational note for Phase 0)
        scope_note_parts = []
        if site_id in sites:
            sa = sites[site_id].get("scope_allocation")
            if sa:
                # Flag if any electricity stream is Scope 3 Cat 13 — resident split needed
                elec_streams = [sa.get("elec_communal_ll"), sa.get("elec_occupied_resi"), sa.get("elec_void")]
                if any(v == "Scope 3 Cat 13" for v in elec_streams if v):
                    scope_note_parts.append(
                        "Resident electricity split pending (Sycous integration). "
                        "Meter-aggregate values recorded; landlord/resident apportionment deferred."
                    )
            else:
                log.append(f"  WARNING: {site_id} missing scope_allocation in sites.json")
        else:
            log.append(f"  WARNING: {site_id} not in sites.json — Arbnco entry recorded without scope context")

        out[site_id] = {
            "id": site_id,
            "meters": {
                "electricity": _to_int(ws.cell(r, 4).value),
                "gas": _to_int(ws.cell(r, 5).value),
            },
            "consumption_kwh": {
                "electricity": _to_float(ws.cell(r, 7).value),
                "gas": _to_float(ws.cell(r, 8).value),
                "total": _to_float(ws.cell(r, 6).value),
            },
            "emissions_actual_tco2e": {
                "electricity": _to_float(ws.cell(r, 10).value),
                "gas": _to_float(ws.cell(r, 11).value),
                "total": _to_float(ws.cell(r, 9).value),
            },
            "emissions_national_avg_tco2e": {
                "electricity": _to_float(ws.cell(r, 13).value),
                "gas": _to_float(ws.cell(r, 14).value),
                "total": _to_float(ws.cell(r, 12).value),
            },
            "source_arbnco_notes": _to_str(ws.cell(r, 16).value),
            "data_period": "CY2025",
            "notes": " | ".join(scope_note_parts) if scope_note_parts else None,
        }
        rows_resolved += 1

    log.append(f"  Raw Data: scanned={rows_scanned}, resolved={rows_resolved}")
    if unmapped_known:
        log.append(f"  Unmapped sites in Arbnco (known, excluded): {unmapped_known}")
    if unmapped_unknown:
        log.append(f"  Unmapped sites in Arbnco (UNKNOWN — investigate): {unmapped_unknown}")
    if portfolio_row is None:
        log.append("  WARNING: portfolio total row not found in Raw Data")

    if portfolio_row is not None:
        out["portfolio"] = {
            **portfolio_row,
            "data_period": "CY2025",
            "sites_with_data": rows_resolved,
        }

    return out
