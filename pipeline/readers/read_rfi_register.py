"""RFI Register reader — produces dist/eir/rfi_status.json.

Reads the RFI_P06 tab from the IVG RFI Register workbook. Header row 12, data row 13+.
Theme column is only set on the first row of each theme group; needs forward-fill.
Status defaults to "Open" in P06 (no status column populated yet).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _to_str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _truncate(s: str | None, limit: int = 200) -> str | None:
    if s is None:
        return None
    if len(s) <= limit:
        return s
    return s[: limit - 1].rstrip() + "…"


def _theme_letter(theme: str | None) -> str | None:
    """Extract leading letter from 'A. Organisational Boundary & Reporting' → 'A'."""
    if theme is None:
        return None
    t = theme.strip()
    if len(t) >= 2 and t[1] == ".":
        return t[0]
    return None


def find_workbook(source_data: Path) -> Path:
    candidates = list(source_data.glob("*SH-1000*RFI*Register*.xlsx"))
    if not candidates:
        # Fallback pattern (some versions omit "RFI" or use different separators)
        candidates = list(source_data.glob("*SH-1000*Register*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No RFI Register workbook (SH-1000*Register*) in {source_data}")
    # Prefer the highest P-revision if multiple
    candidates.sort()
    return candidates[-1]


def read_rfi(workbook_path: Path, log: list[str]) -> dict:
    log.append("")
    log.append(f"[RFI Register] Reading: {workbook_path.name}")

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    if "RFI_P06" not in wb.sheetnames:
        log.append(f"  ERROR: RFI_P06 sheet missing. Available: {wb.sheetnames}")
        return {"items": [], "summary": {"total": 0, "by_theme": {}, "by_status": {}}}

    ws = wb["RFI_P06"]
    items: list[dict] = []
    current_theme: str | None = None

    for r in range(13, ws.max_row + 1):
        ref = _to_str(ws.cell(r, 3).value)
        if ref is None:
            continue
        # Forward-fill theme
        theme_cell = _to_str(ws.cell(r, 2).value)
        if theme_cell is not None:
            current_theme = theme_cell

        sub_theme = _to_str(ws.cell(r, 4).value)
        request = _to_str(ws.cell(r, 5).value)

        items.append({
            "ref": ref,
            "theme": current_theme,
            "theme_letter": _theme_letter(current_theme),
            "sub_theme": sub_theme,
            "request_summary": _truncate(request, 200),
            "status": "Open",
        })

    by_theme: dict[str, int] = {}
    by_status: dict[str, int] = {}
    for item in items:
        tl = item["theme_letter"] or "?"
        by_theme[tl] = by_theme.get(tl, 0) + 1
        by_status[item["status"]] = by_status.get(item["status"], 0) + 1

    log.append(f"  → {len(items)} items, themes: {sorted(by_theme.keys())}, by_theme: {by_theme}")
    return {
        "items": items,
        "summary": {
            "total": len(items),
            "by_theme": by_theme,
            "by_status": by_status,
        },
    }
